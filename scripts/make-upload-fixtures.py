"""가상 자료만 생성하는 파일 읽기 시험용 도구."""
from pathlib import Path
from docx import Document
from openpyxl import Workbook
p=Path('examples/upload-fixtures');p.mkdir(parents=True,exist_ok=True)
d=Document();d.add_paragraph('가상 연습자료 — 실제 현장 근거 아님');d.add_paragraph('객실 상태: 오전 퇴실 객실 20개. 기준 기간은 가상 2026년 10월이다.')
t=d.add_table(rows=2,cols=2);t.cell(0,0).text='항목';t.cell(0,1).text='가상값';t.cell(1,0).text='정비 시간';t.cell(1,1).text='30분'
d.save(p/'synthetic-hotel.docx')
w=Workbook();s=w.active;s.title='가상 운영';s.append(['항목','가상값']);s.append(['객실',20]);s.append(['정비 시간(분)',30]);s.append(['수식 캐시 없음','=B2+B3']);w.save(p/'synthetic-hotel.xlsx')
# 두 페이지 중 2페이지는 텍스트 없는 시험용 페이지. xref를 바이트 위치로 작성한다.
stream=b'BT /F1 12 Tf 60 740 Td (SYNTHETIC HOTEL: rooms 20. Not real evidence.) Tj ET'
objects=[b'<< /Type /Catalog /Pages 2 0 R >>',b'<< /Type /Pages /Kids [3 0 R 6 0 R] /Count 2 >>',b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream',b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << >> >>']
out=bytearray(b'%PDF-1.4\n');offsets=[0]
for i,obj in enumerate(objects,1):offsets.append(len(out));out.extend(str(i).encode()+b' 0 obj\n'+obj+b'\nendobj\n')
xref=len(out);out.extend(b'xref\n0 7\n0000000000 65535 f \n')
for pos in offsets[1:]:out.extend(f'{pos:010d} 00000 n \n'.encode())
out.extend(f'trailer\n<< /Size 7 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode());(p/'synthetic-hotel.pdf').write_bytes(out)
print('가상 PDF·XLSX·DOCX 생성 완료')
